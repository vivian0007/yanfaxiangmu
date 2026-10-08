"""Isolated regression checks; never opens the live database."""
import importlib.util
import os
from pathlib import Path
import secrets
import sys
import sqlite3
import tempfile
import unittest
from io import BytesIO

IMPORT_DATA = tempfile.TemporaryDirectory()
os.environ['VIVIAN_DATA_DIR'] = IMPORT_DATA.name
os.environ['VIVIAN_SECRET'] = secrets.token_hex(32)
os.environ['VIVIAN_ADMIN_PASSWORD'] = secrets.token_urlsafe(20)
spec = importlib.util.spec_from_file_location('pm_app', Path(__file__).parents[1] / 'app.py')
module = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = module
spec.loader.exec_module(module)
module.app.config.update(TESTING=True, SESSION_COOKIE_SECURE=False)

class ProjectManagementTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        module.DATA_DIR = self.tmp.name
        module.DB_PATH = str(Path(self.tmp.name) / 'test.db')
        module.init_db()
        self.db = sqlite3.connect(module.DB_PATH)
        from werkzeug.security import generate_password_hash
        for uid, name, dept, manager in [(2, 'Owner', 'R&D', 0), (3, 'Executor', 'R&D', 0), (4, 'Other', 'Sales', 0), (5, 'Manager', 'R&D', 1)]:
            self.db.execute('INSERT INTO users(id,name,phone,password_hash,department,is_manager) VALUES(?,?,?,?,?,?)',
                            (uid, name, str(uid), generate_password_hash('test-password'), dept, manager))
        self.db.commit()
        self.admin, self.owner, self.executor, self.other, self.manager = [self.client(i) for i in range(1,6)]
        self.pid = self.project(self.owner)

    def tearDown(self):
        self.db.close()
        self.tmp.cleanup()

    def client(self, uid):
        c = module.app.test_client()
        with c.session_transaction() as session:
            session['uid'] = uid
        return c

    def project(self, client, **extra):
        data = dict(project_code='P-001', project_name='Project', category='market', start_date='2026-01-01', delivery_date='2027-01-01',status='planning',priority='high',project_share=100)
        data.update(extra)
        r = client.post('/api/projects', json=data)
        self.assertEqual(r.status_code, 200, r.json)
        return r.json['id']

    def task(self):
        r = self.owner.post(f'/api/projects/{self.pid}/tasks', json=dict(title='Task',assignee_id=3,due_date='2026-10-01',status='todo',priority='high',progress=20))
        self.assertEqual(r.status_code,200,r.json)
        return self.owner.get(f'/api/projects/{self.pid}/tasks').json['tasks'][0]['id']


    def test_manual_weights_permissions_and_persistence(self):
        values = dict(market_weight=42.5, self_weight=17.25)
        self.assertEqual(self.owner.put('/api/settings', json=values).status_code, 403)
        self.assertEqual(self.admin.put('/api/settings', json=values).status_code, 200)
        module.init_db()
        saved = self.admin.get('/api/settings').json
        self.assertEqual(saved['market_weight'], 42.5)
        self.assertEqual(saved['self_weight'], 17.25)
        self.assertEqual(saved['full_score'], 59.75)
        for bad in (-1, 101, 'invalid', None):
            self.assertEqual(self.admin.put('/api/settings', json=dict(market_weight=bad, self_weight=30)).status_code, 400)
        self.assertEqual(self.admin.get('/api/settings').json['market_weight'], 42.5)

    def test_person_weights_scope_and_calculation(self):
        self.db.row_factory = sqlite3.Row
        self.org_fixture()
        values=dict(market_weight=50,self_weight=20)
        self.assertEqual(self.manager.put('/api/person-weights/2',json=values).status_code,200)
        self.assertEqual(self.owner.put('/api/person-weights/3',json=values).status_code,200)
        for uid in (1,2,4,5):
            self.assertEqual(self.owner.put(f'/api/person-weights/{uid}',json=values).status_code,403)
        self.assertEqual(self.executor.put('/api/person-weights/2',json=values).status_code,403)
        self.db.execute('UPDATE users SET is_manager=1,manager_level=3 WHERE id=3')
        self.db.execute("UPDATE users SET department='R&D',department2='Software',department3='Platform' WHERE id=4")
        self.db.commit()
        self.assertEqual(self.executor.put('/api/person-weights/4',json=values).status_code,200)
        self.assertEqual(self.executor.put('/api/person-weights/3',json=values).status_code,403)
        self.db.execute("UPDATE users SET department3='Other' WHERE id=4")
        self.db.commit()
        self.assertEqual(self.executor.put('/api/person-weights/4',json=values).status_code,403)
        self.assertEqual(self.manager.put('/api/person-weights/3',json=values).status_code,200)
        y,m=module.allowed_period()
        self.owner.put(f'/api/projects/{self.pid}/progress',json=dict(year=y,month=m,progress=80))
        perf=module.monthly_performance(self.db,2,y,m)
        self.assertEqual(perf['market_score'],40)
        module.init_db()
        self.assertEqual(module.load_weights(self.db,2),(0.5,0.2))
        self.assertEqual(self.manager.put('/api/person-weights/2',json=dict(market_weight=101,self_weight=20)).status_code,400)
        self.assertEqual(self.manager.put('/api/person-weights/2',json=dict(use_default=True)).status_code,200)
        self.assertEqual(module.load_weights(self.db,2),(0.3,0.3))

    def test_zero_projects_full_category_score(self):
        empty=module.calc_performance([],0.425,0.1725)
        self.assertEqual(empty['total_score'],59.75)
        self.assertEqual(empty['total'],0)
        market=module.calc_performance([dict(category='market',progress=0)],0.4,0.2)
        self.assertEqual(market['market_score'],0)
        self.assertEqual(market['self_score'],20)
        own=module.calc_performance([dict(category='self',progress=50)],0.4,0.2)
        self.assertEqual(own['total_score'],50)
        both=module.calc_performance([dict(category='market',progress=50),dict(category='self',progress=80)],0.4,0.2)
        self.assertEqual(both['total_score'],36)
        self.assertEqual(module.calc_performance([],0,0)['total_score'],0)

    def test_monthly_project_share_limits(self):
        self.assertEqual(self.owner.put(f'/api/projects/{self.pid}/share',json={'share':20}).status_code,200)
        ids=[self.pid]
        for v in (30,40,5):
            ids.append(self.project(self.owner,project_share=v))
        payload=dict(project_code='Extra',project_name='Extra',category='market',start_date='2026-01-01',delivery_date='2027-01-01',project_share=5.01)
        count=self.db.execute('SELECT COUNT(*) FROM projects').fetchone()[0]
        response=self.owner.post('/api/projects',json=payload)
        self.assertEqual(response.status_code,400)
        self.assertIn('5%',response.json['error'])
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM projects').fetchone()[0],count)
        payload['project_share']=5
        last=self.owner.post('/api/projects',json=payload).json['id']
        self.assertEqual(self.owner.put(f'/api/projects/{last}/share',json={'share':5.01}).status_code,400)
        self.assertEqual(self.owner.get(f'/api/projects/{last}/share').json['share'],5)
        self.assertEqual(self.owner.put(f'/api/projects/{ids[0]}/share',json={'share':10}).status_code,200)
        self.assertEqual(self.owner.put(f'/api/projects/{last}/share',json={'share':15}).status_code,200)
        self.project(self.owner,category='self',project_share=100)
        self.project(self.other,project_share=100)
        self.assertEqual(self.other.put(f'/api/projects/{last}/share',json={'share':1}).status_code,403)
        for bad in (-1,101,'NaN',0.001):
            self.assertEqual(self.owner.put(f'/api/projects/{last}/share',json={'share':bad}).status_code,400)
        from unittest.mock import patch
        y,m=module.allowed_period()
        with patch.object(module,'allowed_period',return_value=(y+1,m)):
            self.assertEqual(self.owner.get(f'/api/projects/{last}/share').json['available'],100)
            self.assertEqual(self.owner.get(f'/api/projects/{last}/share').json['share'],0)
            self.assertEqual(self.owner.put(f'/api/projects/{last}/share',json={'share':100}).status_code,200)
        self.assertEqual(self.owner.get(f'/api/projects/{last}/share').json['share'],15)

    def test_weighted_project_performance(self):
        self.owner.put(f'/api/projects/{self.pid}/share',json={'share':20})
        second=self.project(self.owner,project_share=30)
        y,m=module.allowed_period()
        for pid,v in ((self.pid,100),(second,50)):
            self.owner.put(f'/api/projects/{pid}/progress',json=dict(year=y,month=m,progress=v))
        result=self.admin.get(f'/api/performance/2?year={y}').json['months'][str(m)]
        self.assertEqual(result['market_score'],10.5)
        self.assertEqual(result['total_score'],40.5)
        self.assertEqual(self.owner.put(f'/api/projects/{self.pid}',json={'category':'self'}).status_code,400)
        module.init_db()
        self.assertEqual(self.owner.get(f'/api/projects/{self.pid}/share').json['share'],20)

    def test_share_migration_preserves_history(self):
        y,m=module.allowed_period()
        second=self.project(self.owner,project_share=0)
        for pid,v in ((self.pid,80),(second,40)):
            self.owner.put(f'/api/projects/{pid}/progress',json=dict(year=y,month=m,progress=v))
        self.db.execute('DROP TABLE project_shares');self.db.commit()
        module.init_db()
        for pid in (self.pid,second):
            self.assertEqual(self.owner.get(f'/api/projects/{pid}/share').json['share'],50)
        result=self.admin.get(f'/api/performance/2?year={y}').json['months'][str(m)]
        self.assertEqual(result['market_score'],18)
        self.assertEqual(self.owner.put(f'/api/projects/{self.pid}/share',json=dict(share=20,year=y-1,month=m)).status_code,400)
        self.assertEqual(self.owner.get(f'/api/projects/{self.pid}/share').json['share'],50)

    def test_admin_dedicated_password_reset(self):
        route='/api/users/2/reset-password'
        for client in (self.owner,self.manager,module.app.test_client()):
            self.assertIn(client.post(route,json={'password':'temporary-123'}).status_code,(401,403))
        self.assertEqual(self.admin.post(route,json={'password':'short'}).status_code,400)
        self.assertEqual(self.admin.post('/api/users/999/reset-password',json={'password':'temporary-123'}).status_code,404)
        before=self.db.execute('SELECT name,department,is_manager FROM users WHERE id=2').fetchone()
        self.assertEqual(self.admin.post(route,json={'password':'temporary-123'}).status_code,200)
        client=module.app.test_client()
        self.assertFalse(client.post('/api/login',json={'phone':'2','password':'test-password'}).json['ok'])
        self.assertTrue(client.post('/api/login',json={'phone':'2','password':'temporary-123'}).json['ok'])
        self.assertEqual(self.db.execute('SELECT must_change_password FROM users WHERE id=2').fetchone()[0],1)
        self.assertEqual(self.db.execute('SELECT name,department,is_manager FROM users WHERE id=2').fetchone(),before)

    def test_batch_assignment_and_reassignment(self):
        pid=self.project(self.manager)
        response=self.manager.post(f'/api/projects/{pid}/tasks',json=dict(title='Design work',assignee_ids=[2,3]))
        self.assertEqual(response.status_code,200,response.json)
        self.assertEqual(response.json['created'],2)
        tasks=self.manager.get(f'/api/projects/{pid}/tasks').json['tasks']
        self.assertEqual(len(tasks),2)
        for client,uid in ((self.owner,2),(self.executor,3)):
            self.assertIn(pid,[p['id'] for p in client.get('/api/projects').json['projects']])
            self.assertEqual([t['assignee_id'] for t in client.get(f'/api/projects/{pid}/tasks').json['tasks']],[uid])
            self.assertEqual([t['assignee_id'] for t in client.get(f'/api/projects/{pid}').json['tasks']],[uid])
            self.assertEqual(client.post(f'/api/projects/{pid}/tasks',json={'title':'Unauthorized'}).status_code,403)
        a=next(t['id'] for t in tasks if t['assignee_id']==2)
        self.assertEqual(self.executor.put(f'/api/tasks/{a}',json={'progress':50}).status_code,403)
        self.assertEqual(self.owner.put(f'/api/tasks/{a}',json={'assignee_id':3}).status_code,403)
        self.assertEqual(self.manager.put(f'/api/tasks/{a}',json={'assignee_id':3}).status_code,200)
        self.assertNotIn(pid,[p['id'] for p in self.owner.get('/api/projects').json['projects']])
        self.assertEqual(self.owner.get(f'/api/projects/{pid}').status_code,403)
        self.assertEqual(len(self.executor.get(f'/api/projects/{pid}/tasks').json['tasks']),2)
        self.assertEqual(self.other.get(f'/api/projects/{pid}').status_code,403)
        self.assertEqual(self.manager.post(f'/api/projects/{pid}/tasks',json=dict(title='Invalid batch',assignee_ids=[3,999])).status_code,400)
        self.assertEqual(len(self.manager.get(f'/api/projects/{pid}/tasks').json['tasks']),2)

    def test_task_completion_notes(self):
        tid=self.task()
        route=f'/api/tasks/{tid}'
        note='已完成设计和联调。\n待确认验收结果。'
        self.assertEqual(self.executor.put(route,json={'completion_notes':note}).status_code,200)
        task=self.owner.get(f'/api/projects/{self.pid}/tasks').json['tasks'][0]
        self.assertEqual(task['completion_notes'],note)
        self.assertEqual(self.other.put(route,json={'completion_notes':'wrong'}).status_code,403)
        self.assertEqual(self.executor.put(route,json={'completion_notes':123}).status_code,400)
        self.assertEqual(self.executor.put(route,json={'completion_notes':'a'*2001}).status_code,400)
        self.assertEqual(self.executor.put(route,json={'progress':60}).status_code,200)
        module.init_db()
        self.assertEqual(self.owner.get(f'/api/projects/{self.pid}/tasks').json['tasks'][0]['completion_notes'],note)
        self.assertEqual(self.executor.put(route,json={'completion_notes':''}).status_code,200)
        self.assertEqual(self.owner.get(f'/api/projects/{self.pid}/tasks').json['tasks'][0]['completion_notes'],'')

    def test_task_share_atomic_save(self):
        tid=self.task()
        y,m=module.allowed_period()
        second=self.project(self.executor,project_share=95)
        payload=dict(project_share=5,share_year=y,share_month=m,progress=50,completion_notes='Finished design')
        self.assertEqual(self.executor.put(f'/api/tasks/{tid}',json=payload).status_code,200)
        self.assertEqual(self.executor.get(f'/api/projects/{self.pid}/share').json['share'],5)
        payload.update(project_share=5.01,progress=90,completion_notes='Should not save')
        self.assertEqual(self.executor.put(f'/api/tasks/{tid}',json=payload).status_code,400)
        task=self.owner.get(f'/api/projects/{self.pid}/tasks').json['tasks'][0]
        self.assertEqual(task['progress'],50)
        self.assertEqual(task['completion_notes'],'Finished design')
        self.assertEqual(self.executor.get(f'/api/projects/{self.pid}/share').json['share'],5)
        payload.update(project_share=5,share_year=y-1)
        self.assertEqual(self.executor.put(f'/api/tasks/{tid}',json=payload).status_code,400)
        payload['share_year']=y
        self.assertEqual(self.owner.put(f'/api/tasks/{tid}',json=payload).status_code,403)

    def test_task_extension_approval_flow(self):
        from datetime import date,timedelta
        tid=self.task();url=f'/api/tasks/{tid}'
        original=(date.today()+timedelta(days=2)).isoformat()
        later=(date.today()+timedelta(days=5)).isoformat()
        self.assertEqual(self.owner.put(url,json={'due_date':original}).status_code,200)
        self.assertEqual(self.executor.put(url,json={'due_date':later}).status_code,403)
        self.assertEqual(self.other.post(url+'/extensions',json={'requested_date':later,'reason':'Need time'}).status_code,403)
        self.assertEqual(self.executor.post(url+'/extensions',json={'requested_date':original,'reason':'Need time'}).status_code,400)
        self.assertEqual(self.executor.post(url+'/extensions',json={'requested_date':later,'reason':''}).status_code,400)
        self.assertEqual(self.executor.post(url+'/extensions',json={'requested_date':later,'reason':'Waiting for parts'}).status_code,200)
        self.assertEqual(self.executor.post(url+'/extensions',json={'requested_date':later,'reason':'Duplicate'}).status_code,400)
        task=self.owner.get(f'/api/projects/{self.pid}/tasks').json['tasks'][0]
        self.assertEqual(task['due_date'],original)
        eid=task['extension']['id']
        for actor in (self.executor,self.other,self.manager):
            self.assertEqual(actor.put(f'/api/task-extensions/{eid}',json={'decision':'approved'}).status_code,403)
        self.assertEqual(self.owner.put(f'/api/task-extensions/{eid}',json={'decision':'rejected','review_note':'Please revise'}).status_code,200)
        self.assertEqual(self.owner.get(f'/api/projects/{self.pid}/tasks').json['tasks'][0]['due_date'],original)
        self.assertEqual(self.executor.post(url+'/extensions',json={'requested_date':later,'reason':'Revised plan'}).status_code,200)
        eid=self.owner.get(f'/api/projects/{self.pid}/tasks').json['tasks'][0]['extension']['id']
        self.assertEqual(self.owner.put(f'/api/task-extensions/{eid}',json={'decision':'approved'}).status_code,200)
        self.assertEqual(self.executor.get('/api/work/tasks').json['tasks'][0]['due_date'],later)
        self.assertEqual(self.owner.put(f'/api/task-extensions/{eid}',json={'decision':'approved'}).status_code,400)
        next_date=(date.today()+timedelta(days=7)).isoformat()
        self.executor.post(url+'/extensions',json={'requested_date':next_date,'reason':'Another request'})
        eid=self.owner.get(f'/api/projects/{self.pid}/tasks').json['tasks'][0]['extension']['id']
        self.owner.put(url,json={'assignee_id':2})
        self.assertEqual(self.owner.get(f'/api/projects/{self.pid}/tasks').json['tasks'][0]['extension']['status'],'cancelled')
        self.assertEqual(self.owner.put(f'/api/task-extensions/{eid}',json={'decision':'approved'}).status_code,400)

    def test_owner_and_engineer_independent_shares(self):
        tid=self.task();y,m=module.allowed_period()
        self.owner.put(f'/api/projects/{self.pid}/share',json={'share':30})
        payload=dict(project_share=70,share_year=y,share_month=m,completion_notes='My work')
        self.assertEqual(self.executor.put(f'/api/tasks/{tid}',json=payload).status_code,200)
        self.assertEqual(self.owner.get(f'/api/projects/{self.pid}/share').json['share'],30)
        self.assertEqual(self.executor.get(f'/api/projects/{self.pid}/share').json['share'],70)
        payload['project_share']=40
        self.assertEqual(self.executor.put(f'/api/tasks/{tid}',json=payload).status_code,200)
        self.assertEqual(self.owner.get(f'/api/projects/{self.pid}/share').json['share'],30)

    def org_fixture(self):
        self.db.execute("UPDATE users SET department2='Software' WHERE id IN (2,3)")
        self.db.execute("UPDATE users SET department3='Platform' WHERE id=3")
        self.db.execute("UPDATE users SET is_manager=1,manager_level=2 WHERE id=2")
        self.db.commit()

    def test_department_hierarchy_and_export(self):
        self.org_fixture()
        def ids(client, query=''):
            return {u['id'] for u in client.get('/api/dept/members'+query).json['members']}
        self.assertEqual(ids(self.manager),{2,3,5})
        self.assertEqual(ids(self.manager,'?department2=Software'),{2,3})
        self.assertEqual(ids(self.manager,'?scope=direct'),{2})
        self.assertEqual(ids(self.owner),{2,3})
        self.assertEqual(ids(self.owner,'?department2=Other'),set())
        self.assertEqual(ids(self.owner,'?department3=Platform'),{3})
        self.assertEqual(self.owner.get('/api/dept/projects?department2=Other').json['projects'],[])
        self.assertEqual(self.owner.get('/api/dept/performance?department2=Other').json['list'],[])
        self.assertEqual(self.owner.get('/api/dept/export.xlsx?department2=Other').status_code,404)
        self.assertEqual(self.executor.get('/api/dept/members').status_code,403)
        self.assertEqual(self.owner.post('/api/profile',json={'department':'R&D','department2':'Software','position':'Engineer'}).status_code,200)
        response=self.manager.get('/api/dept/export.xlsx?department2=Software')
        self.assertEqual(response.status_code,200)
        wb=module.openpyxl.load_workbook(BytesIO(response.data))
        text=str(list(wb.active.values))
        self.assertIn('Owner',text)
        self.assertNotIn('Other',text)
        response.close()



    def test_existing_rd_department_paths(self):
        self.db.execute("UPDATE users SET department='研发部' WHERE id=5")
        self.db.execute("UPDATE users SET department='研发部-电子电气组',is_manager=1,manager_level=2 WHERE id=2")
        self.db.execute("UPDATE users SET department='研发部-电子电气组-应用组' WHERE id=3")
        self.db.commit()
        people=self.manager.get('/api/dept/members?department2=电子电气组').json['members']
        self.assertEqual({x['id'] for x in people},{2,3})
        self.assertEqual(self.owner.get('/api/dept/members?department3=应用组').json['members'][0]['id'],3)
        self.assertEqual(self.executor.post('/api/profile',json={'department':'研发部','position':'Engineer'}).status_code,200)
        self.assertEqual(self.db.execute('SELECT department FROM users WHERE id=3').fetchone()[0],'研发部')
        self.assertEqual(self.executor.get('/api/me').json['user']['department3'],'应用组')

    def test_third_level_and_sibling_scope(self):
        self.org_fixture()
        self.db.execute("UPDATE users SET is_manager=1, manager_level=3 WHERE id=3")
        # Same third-level name in a different second-level department is not the same scope.
        self.db.execute("UPDATE users SET department='R&D', department2='Hardware', department3='Platform' WHERE id=4")
        self.db.commit()
        people=self.executor.get('/api/dept/members').json['members']
        self.assertEqual([x['id'] for x in people],[3])
        self.assertEqual(self.executor.get('/api/dept/members?department2=Hardware').json['members'],[])
        self.assertEqual([x['id'] for x in self.owner.get('/api/dept/members?scope=direct').json['members']],[3])
        filters=self.manager.get('/api/dept/filters?department2=Unknown').json
        self.assertEqual(filters['departments3'],[])

    def test_org_admin_validation_and_persistence(self):
        base=dict(name='New',phone='100',department='R&D',position='Lead',is_manager=True,
                  department2='Hardware',department3='Circuit',manager_level=2)
        self.assertEqual(self.admin.post('/api/users',json=base).status_code,200)
        uid=self.db.execute("SELECT id FROM users WHERE name='New'").fetchone()[0]
        # Old clients toggling roles must not erase hierarchy fields.
        edit={k:v for k,v in base.items() if k not in ('department2','department3','manager_level')}
        self.assertEqual(self.admin.put(f'/api/users/{uid}',json=edit).status_code,200)
        self.assertEqual(self.db.execute('SELECT department2,department3,manager_level FROM users WHERE id=?',(uid,)).fetchone(),('Hardware','Circuit',2))
        for extra in [dict(department2=''),dict(manager_level=4),dict(manager_level=True)]:
            self.assertEqual(self.admin.put(f'/api/users/{uid}',json={**base,**extra}).status_code,400)
        module.init_db()
        self.assertEqual(self.db.execute('SELECT department2 FROM users WHERE id=?',(uid,)).fetchone()[0],'Hardware')


    def test_first_password_and_regular_change(self):
        self.db.execute('UPDATE users SET must_change_password=1 WHERE id=2')
        self.db.commit()
        c=module.app.test_client()
        self.assertEqual(c.post('/api/login',json={'phone':'2','password':'test-password'}).json['user']['id'],2)
        self.assertEqual(c.get('/api/me').json['user']['id'],2)
        self.assertEqual(c.post('/api/change-password',json={'new_password':'new-test-password'}).status_code,200)
        self.assertFalse(c.get('/api/me').json['user']['must_change_password'])
        self.assertEqual(c.post('/api/change-password',json={'new_password':'next-test-password'}).status_code,400)
        self.assertEqual(c.post('/api/change-password',json={'old_password':'wrong','new_password':'next-test-password'}).status_code,400)
        self.assertEqual(c.post('/api/change-password',json={'old_password':'new-test-password','new_password':'next-test-password'}).status_code,200)
        c.post('/api/logout')
        self.assertEqual(c.get('/api/me').status_code,401)
        self.assertFalse(c.post('/api/login',json={'phone':'2','password':'test-password'}).json['ok'])
        self.assertTrue(c.post('/api/login',json={'phone':'2','password':'next-test-password'}).json['ok'])
        self.assertTrue(self.executor.post('/api/login',json={'phone':'3','password':'test-password'}).json['ok'])


    def test_personal_department_edit_and_roles(self):
        payload=dict(department='研发部',department2='电子电气组',department3='应用组',position='工程师',is_admin=True,is_manager=True,manager_level=1)
        response=self.executor.post('/api/profile',json=payload)
        self.assertEqual(response.status_code,200)
        user=response.json['user']
        self.assertEqual((user['department'],user['department2'],user['department3']),('研发部','电子电气组','应用组'))
        self.assertFalse(user['is_admin']);self.assertFalse(user['is_manager'])
        self.assertEqual(self.executor.post('/api/profile',json={**payload,'department2':''}).status_code,400)
        self.assertEqual(self.executor.get('/api/me').json['user']['department3'],'应用组')


    def test_project_privacy_across_modules_and_hierarchy(self):
        self.org_fixture()
        child=self.project(self.executor)
        outside=self.project(self.other)
        top=self.project(self.manager)
        def ids(c):return {p['id'] for p in c.get('/api/projects').json['projects']}
        self.assertEqual(ids(self.executor),{child})
        self.assertEqual(ids(self.owner),{self.pid,child})
        self.assertEqual(ids(self.manager),{self.pid,child,top})
        self.assertEqual(ids(self.other),{outside})
        self.assertEqual(ids(self.admin),{self.pid,child,outside,top})
        for suffix in ('','/tasks','/milestones','/share'):
            self.assertEqual(self.executor.get(f'/api/projects/{top}'+suffix).status_code,403)
        tid=self.task()
        self.owner.post(f'/api/projects/{self.pid}/milestones',json={'name':'Secret','due_date':'2026-10-01'})
        self.assertEqual([t['id'] for t in self.executor.get('/api/work/tasks').json['tasks']],[tid])
        self.assertEqual([t['id'] for t in self.executor.get('/api/my-tasks').json['tasks']],[tid])
        self.assertEqual(len(self.executor.get('/api/work/milestones').json['milestones']),1)
        self.assertEqual(self.executor.get('/api/dashboard').json['task_total'],1)
        self.assertEqual(self.executor.put(f'/api/tasks/{tid}',json={'progress':90}).status_code,200)
        self.assertEqual(self.manager.post(f'/api/projects/{child}/tasks',json={'title':'No edit'}).status_code,403)
        self.db.execute("UPDATE users SET department2='Elsewhere' WHERE id=3");self.db.commit()
        self.assertNotIn(child,ids(self.owner))
        self.assertIn(child,ids(self.manager))
        self.db.execute("UPDATE users SET department2='Software',department3='Platform',is_manager=1,manager_level=3 WHERE id=3")
        self.db.execute("UPDATE users SET department='R&D',department2='Software',department3='Platform' WHERE id=4");self.db.commit()
        self.assertIn(outside,ids(self.executor))
        self.db.execute("UPDATE users SET department3='Other' WHERE id=4");self.db.commit()
        self.assertNotIn(outside,ids(self.executor))

    def test_shared_monthly_performance_and_history(self):
        pid=self.project(self.manager)
        ids=[]
        for title in ['Work A','Work B']:
            self.manager.post(f'/api/projects/{pid}/tasks',json={'title':title,'assignee_id':5})
        ids=[t['id'] for t in self.manager.get(f'/api/projects/{pid}/tasks').json['tasks']]
        y,m=module.allowed_period()
        for tid,progress in zip(ids,[40,80]):
            r=self.manager.put(f'/api/tasks/{tid}/monthly',json={'year':y,'month':m,'progress':progress,'done_items':'Done'})
            self.assertEqual(r.status_code,200,r.json)
        self.assertEqual(self.manager.put(f'/api/projects/{pid}/share',json={'share':100}).status_code,200)
        result=self.admin.get(f'/api/performance/5?year={y}').json['months'][str(m)]
        self.assertEqual(result['total'],1)
        self.assertEqual(result['market_avg'],60)
        self.assertEqual(result['total_score'],48)  # market 18 + empty self category 30
        self.assertEqual(self.manager.put(f'/api/projects/{pid}/progress',json={'year':y,'month':m,'progress':100}).status_code,200)
        self.assertEqual(self.admin.get(f'/api/performance/5?year={y}').json['months'][str(m)]['total'],1)
        self.assertEqual(self.manager.put(f'/api/tasks/{ids[0]}/monthly',json={'year':y-1,'month':m,'progress':90}).status_code,400)
        self.assertEqual(self.owner.put(f'/api/tasks/{ids[0]}/monthly',json={'year':y,'month':m,'progress':90}).status_code,403)
        self.assertEqual(len(self.manager.get(f'/api/projects/{pid}').json['task_reports']),2)
        self.assertEqual(self.owner.get(f'/api/projects/{pid}').status_code,403)
        self.assertEqual(self.manager.delete(f'/api/tasks/{ids[0]}').status_code,400)
        self.assertEqual(self.manager.delete(f'/api/projects/{pid}').status_code,400)
        self.assertEqual(self.manager.put(f'/api/tasks/{ids[0]}',json={'assignee_id':2}).status_code,200)
        self.assertEqual(self.owner.put(f'/api/tasks/{ids[0]}/monthly',json={'year':y,'month':m,'progress':90}).status_code,400)
        self.assertEqual(self.admin.get(f'/api/performance/5?year={y}').json['months'][str(m)]['market_avg'],60)

    def test_auth_and_static_assets(self):
        c=module.app.test_client()
        self.assertEqual(c.get('/api/work/tasks').status_code,401)
        self.assertEqual(c.post('/api/login',json={'phone':'2','password':'test-password'}).status_code,200)
        self.assertEqual(c.get('/api/me').json['user']['id'],2)
        for path in ['/', '/static/pm.js', '/static/pm.css', '/static/vendor/bootstrap.min.css', '/static/vendor/bootstrap.bundle.min.js']:
            r=c.get(path)
            self.assertEqual(r.status_code,200,path)
            r.close()

    def test_project_owner_and_status(self):
        pid=self.project(self.admin,user_id=3)
        self.assertEqual(self.executor.get(f'/api/projects/{pid}').status_code,200)
        self.assertEqual(self.other.get(f'/api/projects/{self.pid}').status_code,403)
        r=self.owner.put(f'/api/projects/{self.pid}',json={'project_name':'Renamed'})
        self.assertEqual(r.status_code,200,r.json)
        p=self.owner.get(f'/api/projects/{self.pid}').json['project']
        self.assertEqual((p['status'],p['priority']),('planning','high'))
        self.assertEqual(p['project_name'],'Renamed')

    def test_task_lifecycle_and_permissions(self):
        tid=self.task()
        self.assertEqual(len(self.executor.get('/api/work/tasks').json['tasks']),1)
        self.assertEqual(self.other.get('/api/work/tasks').json['tasks'],[])
        self.assertEqual(self.executor.put(f'/api/tasks/{tid}',json={'status':'done'}).status_code,200)
        self.assertEqual(self.owner.put(f'/api/tasks/{tid}',json={'status':'done'}).status_code,200)
        self.assertEqual(self.owner.get(f'/api/projects/{self.pid}/tasks').json['tasks'][0]['progress'],100)
        self.assertEqual(self.executor.put(f'/api/tasks/{tid}',json={'assignee_id':4}).status_code,403)
        self.assertEqual(self.executor.delete(f'/api/tasks/{tid}').status_code,403)
        self.assertEqual(self.other.put(f'/api/tasks/{tid}',json={'status':'done'}).status_code,403)
        self.assertEqual(self.owner.delete(f'/api/tasks/{tid}').status_code,200)

    def test_invalid_task_payload(self):
        for extra in [{'assignee_id':999}, {'due_date':'2026-02-30'},{'progress':101},{'progress':1.5},{'status':'unknown'}]:
            data=dict(title='Task');data.update(extra)
            r=self.owner.post(f'/api/projects/{self.pid}/tasks',json=data)
            self.assertEqual(r.status_code,400,r.json)

    def test_invalid_project_dates(self):
        for extra in [{'start_date':'2026-02-30'}, {'start_date':'2027-02-01','delivery_date':'2027-01-01'}]:
            r=self.owner.put(f'/api/projects/{self.pid}',json=extra)
            self.assertEqual(r.status_code,400,r.json)

    def test_milestone_lifecycle(self):
        url=f'/api/projects/{self.pid}/milestones'
        self.assertEqual(self.owner.post(url,json={'name':'Launch','due_date':'2026-02-30'}).status_code,400)
        self.assertEqual(self.owner.post(url,json={'name':'Launch','due_date':'2027-01-01'}).status_code,200)
        mid=self.owner.get(url).json['milestones'][0]['id']
        self.assertEqual(self.other.get('/api/work/milestones').json['milestones'],[])
        self.assertEqual(self.owner.put(f'/api/milestones/{mid}',json={'name':'Launch','due_date':'2027-01-01','status':'done'}).status_code,200)
        self.assertTrue(self.owner.get(url).json['milestones'][0]['done_at'])
        self.assertEqual(self.other.delete(f'/api/milestones/{mid}').status_code,403)
        self.assertEqual(self.owner.delete(f'/api/milestones/{mid}').status_code,200)

    def test_monthly_performance_is_independent(self):
        year,month=module.allowed_period()
        url=f'/api/projects/{self.pid}/progress'
        self.assertEqual(self.owner.put(url,json={'year':year-1,'month':month,'progress':90}).status_code,400)
        self.assertEqual(self.owner.put(url,json={'year':year,'month':month,'progress':80,'done_items':'Monthly delivery'}).status_code,200)
        tid=self.task()
        self.executor.put(f'/api/tasks/{tid}',json={'status':'done'})
        with module.app.app_context():
            result=module.monthly_performance(module.get_db(),2,year,month)
        self.assertEqual(result['total_score'],54)  # market 24 + empty self category 30
        self.assertEqual(self.db.execute('SELECT progress FROM project_progress WHERE project_id=?',(self.pid,)).fetchone()[0],80)
        r=self.admin.get(f'/api/export/performance.xlsx?year={year}&month={month}')
        self.assertEqual(r.status_code,200)
        import openpyxl
        book=openpyxl.load_workbook(BytesIO(r.data))
        self.assertIn('个人任务月报',book.sheetnames)
        self.assertIn('项目月度占比',book.sheetnames)
        self.assertEqual(len(book.sheetnames),7)
        book.close()
        self.assertEqual(self.owner.get('/api/export/performance.xlsx').status_code,403)
        self.assertEqual(self.manager.get('/api/dept/performance').status_code,200)

    def test_restart_does_not_invent_monthly_reports(self):
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM project_progress').fetchone()[0],0)
        module.init_db();module.init_db()
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM project_progress').fetchone()[0],0)

    def test_preserve_existing_records_and_weights(self):
        self.db.execute('INSERT INTO project_progress(project_id,year,month,progress,done_items) VALUES(?,?,?,?,?)',(self.pid,2025,12,65,'Historical note'))
        self.db.execute("UPDATE settings SET value='0.4' WHERE key='market_weight'")
        self.db.commit()
        before={t:self.db.execute('SELECT * FROM '+t+' ORDER BY 1').fetchall() for t in ['users','projects','project_progress','settings']}
        module.init_db()
        for table,rows in before.items():
            self.assertEqual(self.db.execute('SELECT * FROM '+table+' ORDER BY 1').fetchall(),rows,table)

    def test_delete_project_cascades(self):
        self.task()
        self.owner.post(f'/api/projects/{self.pid}/milestones',json={'name':'Launch'})
        year,month=module.allowed_period()
        self.owner.put(f'/api/projects/{self.pid}/progress',json={'year':year,'month':month,'progress':20})
        self.assertEqual(self.owner.delete(f'/api/projects/{self.pid}').status_code,200)
        for t in ['tasks','milestones','project_progress']:
            self.assertEqual(self.db.execute('SELECT COUNT(*) FROM '+t).fetchone()[0],0)

if __name__=='__main__':
    unittest.main()
